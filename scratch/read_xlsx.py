# -*- coding: utf-8 -*-
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

# Read standard consumption table
wb = openpyxl.load_workbook(r'C:\Users\s 30\Desktop\Korda\Кор АИ\Корда чтение\Расход ткани и наполнителя для стандартных изделий.xlsx', data_only=True)
ws = wb.active
print(f'Sheet: {ws.title}, Rows: {ws.max_row}, Cols: {ws.max_column}')
print('---HEADER---')
for row in ws.iter_rows(min_row=1, max_row=1, max_col=14):
    print([c.value for c in row])
print('---FIRST 25 DATA ROWS---')
for row in ws.iter_rows(min_row=2, max_row=26, max_col=14):
    print([c.value for c in row])

print('\n\n---CLIENT ORDER---')
wb2 = openpyxl.load_workbook(r'C:\Users\s 30\Desktop\Korda\Кор АИ\Корда чтение\Этра 76175_ 76174_ 76178_ 76186_ 76188_ 76187_ 76193_ 76179_ 76177_ 76176_ 76194_ 76190_ 76183_ 76192_ 76180_ 76181_ 76191_ 76182_ 76189_ 76184_ 76185 (1).xlsx', data_only=True)
for sheet in wb2.sheetnames:
    ws2 = wb2[sheet]
    print(f'\nSheet: {sheet}, Rows: {ws2.max_row}, Cols: {ws2.max_column}')
    for row in ws2.iter_rows(min_row=1, max_row=min(30, ws2.max_row), max_col=min(15, ws2.max_column)):
        print([c.value for c in row])
